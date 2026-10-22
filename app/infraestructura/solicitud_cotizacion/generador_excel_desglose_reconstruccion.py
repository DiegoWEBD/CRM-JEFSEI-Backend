from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.dominio.solicitud_cotizacion.solicitud_cotizacion_unidades.desglose_reconstruccion_unidad import DesgloseReconstruccionUnidad
from app.infraestructura.lib.estilo_excel import (
    COLOR_TOTAL,
    FORMATO_MONTO_UF,
    FORMATO_PRORRATEO,
    escribir_celda,
    escribir_encabezado,
    escribir_linea_informativa,
    escribir_titulo,
)

HOJA_DESGLOSE = 'Desglose reconstrucción'

COLUMNAS = (
    'Unidad',
    'Prorrateo alícuota',
    'Prorrateo por monto asegurado',
    'Monto de reconstrucción (UF)',
    'Monto asegurado sugerido (UF)',
)

ETIQUETA_TOTALES = 'TOTAL'

FILA_TITULO = 1
FILA_INFORMACION = 2
FILA_VACIA = 3


class GeneradorExcelDesgloseReconstruccion:
    """Genera el Excel con el desglose de reconstrucción por unidad de una oportunidad.

    Una fila por unidad (prorrateos, monto de reconstrucción y monto asegurado
    sugerido) más una fila de totales.
    """

    def __init__(self, carpeta_base: str = 'documentos/estudios') -> None:
        self.carpeta_base = Path(carpeta_base)

    def generar(
        self,
        id_proceso_comercial: int,
        desglose: list[DesgloseReconstruccionUnidad],
        encabezado: dict | None = None,
    ) -> str:
        nombre = f'desglose_reconstruccion_oportunidad_{id_proceso_comercial}'
        carpeta = self.carpeta_base / nombre
        carpeta.mkdir(parents=True, exist_ok=True)

        libro = Workbook()
        hoja = libro.active
        hoja.title = HOJA_DESGLOSE

        self.__escribir_bloque_encabezado(hoja, encabezado or {}, len(COLUMNAS))

        fila_encabezados = FILA_VACIA + 1
        for columna, texto in enumerate(COLUMNAS, start=1):
            escribir_encabezado(hoja, fila_encabezados, columna, texto)

        fila = fila_encabezados + 1
        for indice, unidad in enumerate(desglose):
            self.__escribir_fila_unidad(hoja, fila + indice, unidad, indice)

        self.__escribir_fila_totales(hoja, fila + len(desglose), desglose)

        hoja.freeze_panes = hoja.cell(row=fila_encabezados + 1, column=1)

        for columna, ancho in enumerate((18, 20, 26, 26, 28), start=1):
            hoja.column_dimensions[get_column_letter(columna)].width = ancho

        ruta = carpeta / f'{nombre}.xlsx'
        libro.save(ruta)

        return str(ruta)

    def __escribir_bloque_encabezado(self, hoja, encabezado: dict, cantidad_columnas: int) -> None:
        escribir_titulo(hoja, FILA_TITULO, 'Desglose de reconstrucción por unidad', cantidad_columnas)

        cliente = encabezado.get('cliente') or '—'
        escribir_linea_informativa(
            hoja,
            FILA_INFORMACION,
            (
                f'Cliente: {cliente}'
                f'   ·   Unidades: {encabezado.get("cantidad_unidades", 0)}'
                f'   ·   Monto asegurado sugerido total: {encabezado.get("monto_asegurado_sugerido_total", 0)} UF'
            ),
            cantidad_columnas,
        )

    def __escribir_fila_unidad(self, hoja, fila: int, unidad: DesgloseReconstruccionUnidad, indice: int) -> None:
        relleno = 'F7F9FC' if indice % 2 == 0 else None

        escribir_celda(hoja, fila, 1, unidad.identificador_unidad, negrita=True, color_relleno=relleno)
        escribir_celda(hoja, fila, 2, unidad.prorrateo, FORMATO_PRORRATEO, color_relleno=relleno)
        escribir_celda(hoja, fila, 3, unidad.prorrateo_pago, FORMATO_PRORRATEO, color_relleno=relleno)
        escribir_celda(hoja, fila, 4, unidad.monto_reconstruccion, FORMATO_MONTO_UF, color_relleno=relleno)
        escribir_celda(hoja, fila, 5, unidad.monto_asegurado_sugerido, FORMATO_MONTO_UF, color_relleno=relleno)

    def __escribir_fila_totales(self, hoja, fila: int, desglose: list[DesgloseReconstruccionUnidad]) -> None:
        escribir_celda(hoja, fila, 1, ETIQUETA_TOTALES, negrita=True, color_relleno=COLOR_TOTAL)
        escribir_celda(
            hoja, fila, 2, sum(u.prorrateo for u in desglose),
            FORMATO_PRORRATEO, negrita=True, color_relleno=COLOR_TOTAL,
        )
        escribir_celda(
            hoja, fila, 3, sum(u.prorrateo_pago for u in desglose),
            FORMATO_PRORRATEO, negrita=True, color_relleno=COLOR_TOTAL,
        )
        escribir_celda(
            hoja, fila, 4, sum(u.monto_reconstruccion for u in desglose),
            FORMATO_MONTO_UF, negrita=True, color_relleno=COLOR_TOTAL,
        )
        escribir_celda(
            hoja, fila, 5, sum(u.monto_asegurado_sugerido for u in desglose),
            FORMATO_MONTO_UF, negrita=True, color_relleno=COLOR_TOTAL,
        )
