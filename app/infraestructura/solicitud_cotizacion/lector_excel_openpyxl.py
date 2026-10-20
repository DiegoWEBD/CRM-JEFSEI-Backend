import io

from openpyxl import load_workbook

from app.dominio.exceptions.dato_invalido_exception import DatoInvalidoException
from app.dominio.solicitud_cotizacion.solicitud_cotizacion_unidades.celda_excel import CeldaExcel
from app.dominio.solicitud_cotizacion.solicitud_cotizacion_unidades.hoja_excel import HojaExcel
from app.dominio.solicitud_cotizacion.solicitud_cotizacion_unidades.lector_alicuotas_excel import (
    LectorAlicuotasExcel,
)

MENSAJE_ARCHIVO_INVALIDO = 'El archivo no se puede abrir como libro de Excel (.xlsx)'


class LectorExcelOpenpyxl(LectorAlicuotasExcel):
    """Lee libros .xlsx con openpyxl, sin ejecutar fórmulas ni macros.

    Se cargan los valores calculados almacenados (`data_only=True`) y, en una
    segunda pasada, las fórmulas (`data_only=False`) para detectar celdas cuya
    fórmula no tiene resultado calculado disponible.
    """

    def obtener_hojas(self, contenido: bytes) -> list[HojaExcel]:
        libro = self._abrir(contenido, data_only=True)

        try:
            return [
                HojaExcel(nombre=nombre, filas=libro[nombre].max_row, columnas=libro[nombre].max_column)
                for nombre in libro.sheetnames
            ]
        finally:
            libro.close()

    def leer_bloque(
        self,
        contenido: bytes,
        hoja: str,
        fila_desde: int,
        fila_hasta: int,
        columnas: list[int],
    ) -> list[list[CeldaExcel]]:
        libro_valores = self._abrir(contenido, data_only=True)
        libro_formulas = self._abrir(contenido, data_only=False)

        try:
            hoja_valores = self._obtener_hoja(libro_valores, hoja)
            hoja_formulas = self._obtener_hoja(libro_formulas, hoja)

            filas: list[list[CeldaExcel]] = []

            for numero_fila in range(fila_desde, fila_hasta + 1):
                fila: list[CeldaExcel] = []

                for numero_columna in columnas:
                    celda_valor = hoja_valores.cell(row=numero_fila, column=numero_columna)
                    celda_formula = hoja_formulas.cell(row=numero_fila, column=numero_columna)

                    formula_sin_resultado = (
                        celda_valor.value is None
                        and getattr(celda_formula, 'data_type', None) == 'f'
                    )

                    fila.append(CeldaExcel(
                        valor=celda_valor.value,
                        formato_numerico=celda_valor.number_format,
                        formula_sin_resultado=formula_sin_resultado,
                    ))

                filas.append(fila)

            return filas
        finally:
            libro_valores.close()
            libro_formulas.close()

    def _abrir(self, contenido: bytes, data_only: bool):
        try:
            return load_workbook(io.BytesIO(contenido), data_only=data_only)
        except Exception as exc:
            raise DatoInvalidoException(MENSAJE_ARCHIVO_INVALIDO) from exc

    def _obtener_hoja(self, libro, nombre: str):
        if nombre not in libro.sheetnames:
            raise DatoInvalidoException(f'La hoja "{nombre}" no existe en el archivo')

        return libro[nombre]
