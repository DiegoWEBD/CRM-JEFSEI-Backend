from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.dominio.cotizacion.cotizacion import Cotizacion
from app.dominio.estudio_comercial.desglose_unidades.desglose_unidad import DesgloseUnidad
from app.dominio.estudio_comercial.desglose_unidades.totales_prima import TotalesPrima


FORMATO_MONTO = '#,##0.00'
FORMATO_PRORRATEO = '0.000000000'

HOJA_DESGLOSE = 'Desglose por unidad'

FILA_GRUPOS = 1
FILA_ENCABEZADOS = 2
PRIMERA_FILA_UNIDADES = 3

# Columnas compartidas por todas las cotizaciones
COLUMNAS_COMPARTIDAS = (
    'Unidad',
    'Prorrateo alícuota',
    'Prorrateo por monto asegurado',
    'Monto de reconstrucción (UF)',
    'Monto asegurado sugerido (UF)',
)

# Columnas que se repiten para cada cotización
COLUMNAS_COTIZACION = (
    'Prima afecta (UF)',
    'Prima excenta (UF)',
    'Prima neta (UF)',
    'Prima IVA (UF)',
    'Prima bruta (UF)',
)

COLORES_GRUPO = ('FFF2CC', 'DDEBF7', 'FCE4D6', 'E2EFDA', 'EDEDED')

ETIQUETA_TOTALES = 'TOTAL'
ETIQUETA_REFERENCIA = 'TOTALES COTIZACIÓN (referencia)'

# (cotizacion, totales de referencia, desglose por unidad) de una cotización
DetalleDesglose = tuple[Cotizacion, TotalesPrima, list[DesgloseUnidad]]


class GeneradorExcelDesgloseUnidades:
    """Genera el Excel con el desglose por unidad de las cotizaciones de una solicitud.

    Hoja única con el desglose completo: por cada unidad, su prorrateo de alícuota,
    el prorrateo por monto asegurado (el que reparte las primas), el monto de
    reconstrucción y el monto asegurado sugerido, y por cada cotización las primas
    repartidas según ese prorrateo. Cierra con una fila de totales y una fila de
    referencia con los totales de cada cotización.
    """

    def __init__(self, carpeta_base: str = 'documentos/estudios') -> None:
        self.carpeta_base = Path(carpeta_base)

    def generar(
        self,
        id_solicitud: int,
        datos: list[DetalleDesglose],
    ) -> str:
        nombre = f'desglose_unidades_solicitud_{id_solicitud}'
        carpeta = self.carpeta_base / nombre
        carpeta.mkdir(parents=True, exist_ok=True)

        libro = Workbook()
        hoja = libro.active
        hoja.title = HOJA_DESGLOSE

        self.__escribir_grupos(hoja, [cotizacion.company.nombre.upper() for cotizacion, _, _ in datos])
        self.__escribir_encabezados(hoja, datos)

        unidades = datos[0][2] if datos else []

        for indice, unidad in enumerate(unidades):
            filas_cotizacion = [desglose[indice] for _, _, desglose in datos]
            self.__escribir_fila_unidad(hoja, PRIMERA_FILA_UNIDADES + indice, unidad, filas_cotizacion)

        fila_totales = PRIMERA_FILA_UNIDADES + len(unidades)
        self.__escribir_fila_totales(hoja, fila_totales, datos)
        self.__escribir_fila_referencia(hoja, fila_totales + 1, datos)

        ruta = carpeta / f'{nombre}.xlsx'
        libro.save(ruta)

        return str(ruta)

    def __escribir_grupos(self, hoja, nombres_companies: list[str]) -> None:
        cantidad_compartidas = len(COLUMNAS_COMPARTIDAS)

        celda = hoja.cell(row=FILA_GRUPOS, column=1, value='Unidades')
        hoja.merge_cells(
            start_row=FILA_GRUPOS, start_column=1, end_row=FILA_GRUPOS, end_column=cantidad_compartidas
        )
        self.__estilo_grupo(celda, COLORES_GRUPO[0])

        for indice, nombre in enumerate(nombres_companies):
            inicio = cantidad_compartidas + 1 + indice * len(COLUMNAS_COTIZACION)
            fin = inicio + len(COLUMNAS_COTIZACION) - 1

            celda = hoja.cell(row=FILA_GRUPOS, column=inicio, value=nombre)
            hoja.merge_cells(
                start_row=FILA_GRUPOS, start_column=inicio, end_row=FILA_GRUPOS, end_column=fin
            )
            self.__estilo_grupo(celda, COLORES_GRUPO[(indice + 1) % len(COLORES_GRUPO)])

    def __escribir_encabezados(self, hoja, datos: list[DetalleDesglose]) -> None:
        encabezados = list(COLUMNAS_COMPARTIDAS)

        for _ in datos:
            encabezados.extend(COLUMNAS_COTIZACION)

        for columna, encabezado in enumerate(encabezados, start=1):
            celda = hoja.cell(row=FILA_ENCABEZADOS, column=columna, value=encabezado)
            celda.font = Font(bold=True)
            celda.alignment = Alignment(horizontal='center')

        for columna in range(1, len(encabezados) + 1):
            hoja.column_dimensions[get_column_letter(columna)].width = 20

    def __escribir_fila_unidad(
        self,
        hoja,
        fila: int,
        unidad: DesgloseUnidad,
        filas_cotizacion: list[DesgloseUnidad],
    ) -> None:
        valores = [
            unidad.identificador_unidad,
            unidad.prorrateo,
            unidad.prorrateo_pago,
            unidad.monto_reconstruccion,
            unidad.monto_asegurado_sugerido,
        ]
        formatos = [None, FORMATO_PRORRATEO, FORMATO_PRORRATEO, FORMATO_MONTO, FORMATO_MONTO]

        for fila_cotizacion in filas_cotizacion:
            valores.extend([
                fila_cotizacion.prima_afecta,
                fila_cotizacion.prima_excenta,
                fila_cotizacion.prima_neta,
                fila_cotizacion.prima_iva,
                fila_cotizacion.prima_bruta,
            ])
            formatos.extend([FORMATO_MONTO] * len(COLUMNAS_COTIZACION))

        self.__escribir_fila(hoja, fila, valores, formatos)

    def __escribir_fila_totales(self, hoja, fila: int, datos: list[DetalleDesglose]) -> None:
        unidades = datos[0][2] if datos else []

        valores = [
            ETIQUETA_TOTALES,
            round(sum(unidad.prorrateo or 0.0 for unidad in unidades), 9),
            round(sum(unidad.prorrateo_pago or 0.0 for unidad in unidades), 9),
            round(sum(unidad.monto_reconstruccion or 0.0 for unidad in unidades), 2),
            round(sum(unidad.monto_asegurado_sugerido or 0.0 for unidad in unidades), 2),
        ]
        formatos = [None, FORMATO_PRORRATEO, FORMATO_PRORRATEO, FORMATO_MONTO, FORMATO_MONTO]

        for _, _, desglose in datos:
            valores.extend([
                round(sum(unidad.prima_afecta for unidad in desglose), 2),
                round(sum(unidad.prima_excenta for unidad in desglose), 2),
                round(sum(unidad.prima_neta for unidad in desglose), 2),
                round(sum(unidad.prima_iva for unidad in desglose), 2),
                round(sum(unidad.prima_bruta for unidad in desglose), 2),
            ])
            formatos.extend([FORMATO_MONTO] * len(COLUMNAS_COTIZACION))

        self.__escribir_fila(hoja, fila, valores, formatos, negrita=True)

    def __escribir_fila_referencia(self, hoja, fila: int, datos: list[DetalleDesglose]) -> None:
        valores = [ETIQUETA_REFERENCIA, None, None, None, None]
        formatos = [None, None, None, None, None]

        for _, totales, _ in datos:
            valores.extend([
                totales.prima_afecta,
                totales.prima_excenta,
                totales.prima_neta,
                totales.prima_iva,
                totales.prima_bruta,
            ])
            formatos.extend([FORMATO_MONTO] * len(COLUMNAS_COTIZACION))

        self.__escribir_fila(hoja, fila, valores, formatos, negrita=True)

    def __escribir_fila(
        self,
        hoja,
        fila: int,
        valores: list,
        formatos: list,
        negrita: bool = False,
    ) -> None:
        for columna, (valor, formato) in enumerate(zip(valores, formatos), start=1):
            celda = hoja.cell(row=fila, column=columna, value=valor)

            if negrita:
                celda.font = Font(bold=True)

            if formato is not None:
                celda.number_format = formato

    def __estilo_grupo(self, celda, color: str) -> None:
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal='center')
        celda.fill = PatternFill(fill_type='solid', fgColor=color)
