from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Formatos numéricos compartidos por los Excel de desglose
FORMATO_MONTO_UF = '#,##0.00'
FORMATO_MONTO_CLP = '$ #,##0'
FORMATO_PRORRATEO = '0.000000000'

# Paleta de colores
COLOR_TITULO = '1F4E78'
COLOR_ENCABEZADO = 'D9E2F3'
COLOR_TOTAL = 'FFF2CC'
COLOR_REFERENCIA = 'EDEDED'
COLORES_GRUPO = ('FFF2CC', 'DDEBF7', 'FCE4D6', 'E2EFDA', 'EDEDED')
COLOR_BORDE = 'BFBFBF'

BORDE = Border(
    left=Side(style='thin', color=COLOR_BORDE),
    right=Side(style='thin', color=COLOR_BORDE),
    top=Side(style='thin', color=COLOR_BORDE),
    bottom=Side(style='thin', color=COLOR_BORDE),
)


def escribir_titulo(hoja, fila: int, texto: str, cantidad_columnas: int) -> None:
    celda = hoja.cell(row=fila, column=1, value=texto)
    hoja.merge_cells(
        start_row=fila, start_column=1, end_row=fila, end_column=cantidad_columnas
    )
    celda.font = Font(bold=True, size=14, color='FFFFFF')
    celda.alignment = Alignment(horizontal='center', vertical='center')
    celda.fill = PatternFill(fill_type='solid', fgColor=COLOR_TITULO)
    hoja.row_dimensions[fila].height = 26


def escribir_linea_informativa(hoja, fila: int, texto: str, cantidad_columnas: int) -> None:
    celda = hoja.cell(row=fila, column=1, value=texto)
    hoja.merge_cells(
        start_row=fila, start_column=1, end_row=fila, end_column=cantidad_columnas
    )
    celda.font = Font(italic=True, size=10, color='404040')
    celda.alignment = Alignment(horizontal='left', vertical='center')
    hoja.row_dimensions[fila].height = 16


def escribir_grupo(hoja, fila: int, inicio: int, fin: int, texto: str, color: str) -> None:
    celda = hoja.cell(row=fila, column=inicio, value=texto)
    if fin > inicio:
        hoja.merge_cells(
            start_row=fila, start_column=inicio, end_row=fila, end_column=fin
        )
    celda.font = Font(bold=True, size=10)
    celda.alignment = Alignment(horizontal='center', vertical='center')
    celda.fill = PatternFill(fill_type='solid', fgColor=color)

    for columna in range(inicio, fin + 1):
        hoja.cell(row=fila, column=columna).border = BORDE


def escribir_encabezado(hoja, fila: int, columna: int, texto: str) -> None:
    celda = hoja.cell(row=fila, column=columna, value=texto)
    celda.font = Font(bold=True, size=10)
    celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    celda.fill = PatternFill(fill_type='solid', fgColor=COLOR_ENCABEZADO)
    celda.border = BORDE


def escribir_celda(
    hoja,
    fila: int,
    columna: int,
    valor,
    formato: str | None = None,
    negrita: bool = False,
    color_relleno: str | None = None,
    alineacion: str | None = None,
) -> None:
    celda = hoja.cell(row=fila, column=columna, value=valor)
    celda.border = BORDE

    if negrita:
        celda.font = Font(bold=True, size=10)
    else:
        celda.font = Font(size=10)

    if formato is not None:
        celda.number_format = formato

    if color_relleno is not None:
        celda.fill = PatternFill(fill_type='solid', fgColor=color_relleno)

    if alineacion is not None:
        celda.alignment = Alignment(horizontal=alineacion, vertical='center')
