import io

import pytest
from openpyxl import Workbook

from app.dominio.exceptions.dato_invalido_exception import DatoInvalidoException
from app.infraestructura.solicitud_cotizacion.lector_excel_openpyxl import LectorExcelOpenpyxl


def _libro(hojas: dict) -> bytes:
    """Genera un .xlsx sintético en memoria: {nombre_hoja: lista de filas}."""

    wb = Workbook()
    wb.remove(wb.active)

    for nombre, filas in hojas.items():
        hoja = wb.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.mark.unit
class TestLectorExcelOpenpyxl:

    def test_obtener_hojas(self):
        contenido = _libro({'Alícuotas': [['a', 'b'], ['c', 'd']], 'Resumen': [['x']]})

        hojas = LectorExcelOpenpyxl().obtener_hojas(contenido)

        assert [h.nombre for h in hojas] == ['Alícuotas', 'Resumen']
        assert hojas[0].filas == 2
        assert hojas[0].columnas == 2

    def test_leer_bloque_respeta_orden_de_columnas(self):
        contenido = _libro({'Hoja1': [['Depto 1', 5, 'x'], ['Depto 2', 6, 'y']]})

        filas = LectorExcelOpenpyxl().leer_bloque(contenido, 'Hoja1', 1, 2, [3, 1])

        assert len(filas) == 2
        assert filas[0][0].valor == 'x'
        assert filas[0][1].valor == 'Depto 1'
        assert filas[1][1].valor == 'Depto 2'

    def test_expone_formato_numerico_de_porcentaje(self):
        contenido = _libro_con_formato_porcentaje()

        filas = LectorExcelOpenpyxl().leer_bloque(contenido, 'Hoja1', 1, 1, [2])
        assert filas[0][0].valor == 0.0085
        assert '%' in filas[0][0].formato_numerico

    def test_formula_sin_resultado_calculado(self):
        contenido = _libro_con_formula()

        filas = LectorExcelOpenpyxl().leer_bloque(contenido, 'Hoja1', 2, 2, [2])

        assert filas[0][0].valor is None
        assert filas[0][0].formula_sin_resultado is True

    def test_identificadores_de_texto_conservan_ceros_iniciales(self):
        contenido = _libro({'Hoja1': [['007', 5]]})

        filas = LectorExcelOpenpyxl().leer_bloque(contenido, 'Hoja1', 1, 1, [1])

        assert filas[0][0].valor == '007'

    def test_contenido_no_legible_es_rechazado(self):
        with pytest.raises(DatoInvalidoException):
            LectorExcelOpenpyxl().obtener_hojas(b'esto no es un xlsx')

    def test_hoja_inexistente_es_rechazada(self):
        contenido = _libro({'Hoja1': [['a']]})

        with pytest.raises(DatoInvalidoException):
            LectorExcelOpenpyxl().leer_bloque(contenido, 'NoExiste', 1, 1, [1])


def _libro_con_formato_porcentaje() -> bytes:
    wb = Workbook()
    hoja = wb.active
    hoja.title = 'Hoja1'
    hoja.append(['Depto 1', 0.0085])
    hoja['B1'].number_format = '0.00%'
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _libro_con_formula() -> bytes:
    wb = Workbook()
    hoja = wb.active
    hoja.title = 'Hoja1'
    hoja.append(['Depto 1', 10])
    hoja.append(['Depto 2', None])
    hoja['B2'] = '=B1*2'
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
